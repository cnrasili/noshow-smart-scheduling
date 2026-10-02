"""Train logistic regression and random forest on the cleaned data and write three Excel reports
to data/processed/: logistic_regression.xlsx, random_forest.xlsx, auc_calibration.xlsx.

Split: by appointment date (older appointments train, the latest ~20 % of rows test).
Features follow docs/features.md (SMS_received is excluded).
Rule baseline R4 (docs/formulas/rules.md) uses gender as the group in place of department.
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L
from sklearn.ensemble import RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from scipy.optimize import brentq
from sklearn.inspection import permutation_importance
from sklearn.metrics import brier_score_loss, log_loss, roc_auc_score
from sklearn.preprocessing import StandardScaler

DATA = Path(__file__).parent / "data" / "processed"
SRC = DATA / "appointments_clean.xlsx"
FEATURES = ["LeadDays", "Weekday", "Age", "GenderMale", "Scholarship", "Hypertension", "Diabetes",
            "Alcoholism", "Handicap", "PriorApptCount", "PriorNoShowCount"]
K, K_P, EPS = 5, 5, 0.001
BOLD = Font(bold=True)


def logit(x):
    x = np.clip(x, EPS, 1 - EPS)
    return np.log(x / (1 - x))


def load():
    df = pd.read_excel(SRC, sheet_name="data")
    df["GenderMale"] = (df["Gender"] == "M").astype(int)
    dates = np.sort(df["AppointmentDay"].dt.normalize().unique())
    counts = df.groupby(df["AppointmentDay"].dt.normalize()).size().reindex(dates).cumsum() / len(df)
    cut = counts[counts >= 0.80].index[0]  # last day of the training period
    train = df[df["AppointmentDay"].dt.normalize() <= cut].copy()
    test = df[df["AppointmentDay"].dt.normalize() > cut].copy()
    return df, train, test, cut


def r4_predict(train, test):
    """Rule R4 with gender as the context group (rules.md). Returns no-show probability."""
    s_bar = train["ShowUp"].mean()

    def smoothed(col):
        g = train.groupby(col)["ShowUp"].agg(["sum", "count"])
        return ((g["sum"] + K * s_bar) / (g["count"] + K)).to_dict()

    s_gender, s_age = smoothed("Gender"), smoothed("AgeBand")
    sg = test["Gender"].map(s_gender).fillna(s_bar).to_numpy()
    sa = test["AgeBand"].astype(str).map(s_age).fillna(s_bar).to_numpy()
    z0 = logit(s_bar) + (logit(sg) - logit(s_bar)) + (logit(sa) - logit(s_bar))
    pi0 = 1 / (1 + np.exp(-z0))
    n = test["PriorApptCount"].to_numpy()
    shows = n - test["PriorNoShowCount"].to_numpy()
    return 1 - (shows + K_P * pi0) / (n + K_P), {"s_bar": s_bar, "gender": s_gender, "age": s_age}


def metrics(y, p):
    return {"AUC": roc_auc_score(y, p), "Brier": brier_score_loss(y, p), "LogLoss": log_loss(y, p)}


def calibration_stats(y, p):
    """Calibration-in-the-large (intercept with slope fixed at 1) and calibration slope."""
    z = logit(p)
    a = brentq(lambda a: np.mean(1 / (1 + np.exp(-(z + a)))) - np.mean(y), -5, 5)
    slope = LogisticRegression(C=1e6).fit(z.reshape(-1, 1), y).coef_[0][0]
    return a, slope


def bootstrap_auc_ci(y, p, n=200, seed=0):
    rng = np.random.default_rng(seed)
    y, p = np.asarray(y), np.asarray(p)
    aucs = [roc_auc_score(y[i], p[i]) for i in (rng.integers(0, len(y), len(y)) for _ in range(n))]
    return np.percentile(aucs, [2.5, 97.5])


def header(ws, row, values):
    for i, v in enumerate(values, 1):
        ws.cell(row, i, v).font = BOLD


def dump_df(ws, df, start=1):
    header(ws, start, list(df.columns))
    for r, row in enumerate(df.itertuples(index=False), start + 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v.item() if hasattr(v, "item") else v)
    ws.freeze_panes = ws.cell(start + 1, 1)
    for c in range(1, len(df.columns) + 1):
        ws.column_dimensions[L(c)].width = 18


def info_sheet(wb, rows):
    ws = wb.create_sheet("summary", 0)
    for r, (k, v) in enumerate(rows, 1):
        ws.cell(r, 1, k).font = BOLD
        ws.cell(r, 2, v)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60


def split_info(train, test, cut):
    return [("Train rows", len(train)), ("Test rows", len(test)),
            ("Train period ends", f"{cut:%Y-%m-%d}"),
            ("Test period", f"{test.AppointmentDay.min():%Y-%m-%d} to {test.AppointmentDay.max():%Y-%m-%d}"),
            ("Target", "NoShow (1 = missed the appointment); predictions are no-show probabilities"),
            ("Features", ", ".join(FEATURES)), ("Excluded", "SMS_received (sent after booking)")]


def logistic_report(train, test, cut):
    sc = StandardScaler().fit(train[FEATURES])
    lr = LogisticRegression(max_iter=1000).fit(sc.transform(train[FEATURES]), train["NoShow"])
    coef = lr.coef_[0] / sc.scale_
    b0 = lr.intercept_[0] - float(np.sum(lr.coef_[0] * sc.mean_ / sc.scale_))
    z = test[FEATURES].to_numpy() @ coef + b0
    p = 1 / (1 + np.exp(-z))
    assert np.allclose(p, lr.predict_proba(sc.transform(test[FEATURES]))[:, 1])

    wb = Workbook()
    info_sheet(wb, split_info(train, test, cut) + [("Model", "Logistic regression (L2-penalised, sklearn default C=1, lbfgs, unweighted; features standardised for fitting, coefficients converted to raw scale)")] +
               [(f"Test {k}", round(v, 4)) for k, v in metrics(test["NoShow"], p).items()])
    ws = wb.create_sheet("coefficients")
    header(ws, 1, ["Term", "Coefficient", "Odds ratio"])
    ws.cell(2, 1, "Intercept"); ws.cell(2, 2, float(b0)); ws.cell(2, 3, "")
    for i, (f, c) in enumerate(zip(FEATURES, coef), 3):
        ws.cell(i, 1, f); ws.cell(i, 2, float(c)); ws.cell(i, 3, f"=EXP(B{i})")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16

    ws = wb.create_sheet("predictions")
    cols = ["AppointmentID"] + FEATURES + ["NoShow", "p_python", "p_excel_formula"]
    header(ws, 1, cols)
    for r, (row, pp) in enumerate(zip(test[["AppointmentID"] + FEATURES + ["NoShow"]].itertuples(index=False), p), 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v.item() if hasattr(v, "item") else v)
        nf = len(FEATURES)
        ws.cell(r, nf + 3, float(pp))
        terms = "+".join(f"coefficients!$B${i + 3}*{L(i + 2)}{r}" for i in range(nf))
        ws.cell(r, nf + 4, f"=1/(1+EXP(-(coefficients!$B$2+{terms})))")
    ws.freeze_panes = "B2"
    wb.save(DATA / "logistic_regression.xlsx")
    return p


def forest_report(train, test, cut):
    rf = RandomForestClassifier(n_estimators=300, min_samples_leaf=50, n_jobs=-1, random_state=42)
    rf.fit(train[FEATURES], train["NoShow"])
    p = rf.predict_proba(test[FEATURES])[:, 1]
    wb = Workbook()
    info_sheet(wb, split_info(train, test, cut) +
               [("Model", "Random forest: 300 trees, min_samples_leaf=50, random_state=42, unweighted")] +
               [(f"Test {k}", round(v, 4)) for k, v in metrics(test["NoShow"], p).items()])
    imp = pd.DataFrame({"Feature": FEATURES, "Importance": rf.feature_importances_}).sort_values(
        "Importance", ascending=False)
    dump_df(wb.create_sheet("feature_importance"), imp)
    pi = permutation_importance(rf, test[FEATURES], test["NoShow"], scoring="roc_auc", n_repeats=5,
                                random_state=0, n_jobs=-1)
    pimp = pd.DataFrame({"Feature": FEATURES, "AUC drop when shuffled (test)": pi.importances_mean}).sort_values(
        "AUC drop when shuffled (test)", ascending=False)
    dump_df(wb.create_sheet("permutation_importance"), pimp)
    out = test[["AppointmentID"] + FEATURES + ["NoShow"]].copy()
    out["p_python"] = p
    dump_df(wb.create_sheet("predictions"), out)
    wb.save(DATA / "random_forest.xlsx")
    return p


def eval_block(ws, name, pcol, n, top):
    """ROC table (threshold grid) and calibration table (deciles) as live Excel formulas."""
    last = n + 1
    y, p = f"predictions!$B$2:$B${last}", f"predictions!${pcol}$2:${pcol}${last}"
    ws.cell(top, 1, name).font = Font(bold=True, size=13)
    ws.cell(top + 1, 1, "Positives (no-show)"); ws.cell(top + 1, 2, f"=SUM({y})")
    ws.cell(top + 2, 1, "Negatives (show)"); ws.cell(top + 2, 2, f"=COUNT({y})-B{top + 1}")
    ws.cell(top + 3, 1, "Brier score"); ws.cell(top + 3, 2, f"=SUMPRODUCT(({p}-{y})^2)/COUNT({y})")
    ws.cell(top + 4, 1, "AUC (trapezoid, 0.005 grid)")
    r0 = top + 7
    header(ws, r0 - 1, ["Threshold", "TPR", "FPR"])
    steps = 201
    for i in range(steps):
        r = r0 + i
        ws.cell(r, 1, round(1 - i * 0.005, 3))
        ws.cell(r, 2, f'=COUNTIFS({y},1,{p},">="&A{r})/$B${top + 1}')
        ws.cell(r, 3, f'=COUNTIFS({y},0,{p},">="&A{r})/$B${top + 2}')
    rl = r0 + steps - 1
    ws.cell(top + 4, 2, f"=SUMPRODUCT((C{r0 + 1}:C{rl}-C{r0}:C{rl - 1})*(B{r0 + 1}:B{rl}+B{r0}:B{rl - 1}))/2")

    cc = 6  # calibration table starts at column F
    header_row = top + 6
    for i, h in enumerate(["Decile", "Lower", "Upper", "Count", "Mean predicted", "Observed rate"]):
        ws.cell(header_row, cc + i, h).font = BOLD
    for d in range(10):
        r = header_row + 1 + d
        ws.cell(r, cc, d + 1)
        ws.cell(r, cc + 1, f"=PERCENTILE({p},{d}/10)")
        ws.cell(r, cc + 2, f"=PERCENTILE({p},{d + 1}/10)")
        lo, hi = f"{L(cc + 1)}{r}", f"{L(cc + 2)}{r}"
        cond = f'{p},">="&{lo},{p},' + ('"<="&' if d == 9 else '"<"&') + hi
        ws.cell(r, cc + 3, f"=COUNTIFS({cond})")
        ws.cell(r, cc + 4, f"=AVERAGEIFS({p},{cond})")
        ws.cell(r, cc + 5, f"=AVERAGEIFS({y},{cond})")

    roc = ScatterChart(); roc.title = f"ROC - {name}"; roc.style = 13
    roc.x_axis.title, roc.y_axis.title = "FPR", "TPR"
    s = Series(Reference(ws, min_col=2, min_row=r0, max_row=rl), Reference(ws, min_col=3, min_row=r0, max_row=rl), title="ROC")
    s.marker.symbol = "none"; roc.series.append(s)
    roc.height, roc.width = 7.5, 9
    ws.add_chart(roc, f"{L(cc + 7)}{top}")
    cal = ScatterChart(); cal.title = f"Calibration - {name}"; cal.style = 13
    cal.x_axis.title, cal.y_axis.title = "Mean predicted", "Observed rate"
    rr = Reference(ws, min_col=cc + 5, min_row=header_row + 1, max_row=header_row + 10)
    xx = Reference(ws, min_col=cc + 4, min_row=header_row + 1, max_row=header_row + 10)
    s = Series(rr, xx, title="Model"); s.marker.symbol = "circle"; cal.series.append(s)
    s = Series(xx, xx, title="Perfect"); s.marker.symbol = "none"; cal.series.append(s)
    cal.height, cal.width = 7.5, 9
    ws.add_chart(cal, f"{L(cc + 7 + 10)}{top}")
    return rl + 3


def comparison_report(train, test, cut, preds, r4info):
    n = len(test)
    wb = Workbook()
    y = test["NoShow"]
    rows = split_info(train, test, cut) + [("R4 rule", "docs/formulas/rules.md; gender used as group instead of department")]
    for name, p in preds.items():
        lo, hi = bootstrap_auc_ci(y, p)
        a, sl = calibration_stats(y, p)
        rows += [(f"{name} AUC (exact, Python)", round(roc_auc_score(y, p), 4)),
                 (f"{name} AUC 95% CI (bootstrap, 200 resamples)", f"{lo:.3f} - {hi:.3f}"),
                 (f"{name} Brier (Python)", round(brier_score_loss(y, p), 4)),
                 (f"{name} calibration-in-the-large (0 = ideal)", round(a, 3)),
                 (f"{name} calibration slope (1 = ideal)", round(sl, 3))]
    base = np.full(len(y), train["NoShow"].mean())
    rows += [("No-show rate: train / test", f"{train['NoShow'].mean():.4f} / {y.mean():.4f}"),
             ("Brier of constant predictor (train rate)", round(brier_score_loss(y, base), 4)),
             ("Note", "All models over-predict slightly because the no-show rate fell from train to test (prevalence shift)"),
             ("How to read", "evaluation sheet: ROC and calibration tables are Excel formulas over the predictions sheet")]
    info_sheet(wb, rows)
    ws = wb.create_sheet("predictions")
    header(ws, 1, ["AppointmentID", "NoShow"] + [f"p_{k}" for k in preds] + ["Gender", "AgeBand", "PriorApptCount"])
    for r in range(n):
        ws.cell(r + 2, 1, int(test["AppointmentID"].iloc[r])); ws.cell(r + 2, 2, int(y.iloc[r]))
        for c, p in enumerate(preds.values(), 3):
            ws.cell(r + 2, c, float(p[r]))
        ws.cell(r + 2, 3 + len(preds), test["Gender"].iloc[r])
        ws.cell(r + 2, 4 + len(preds), str(test["AgeBand"].iloc[r]))
        ws.cell(r + 2, 5 + len(preds), int(test["PriorApptCount"].iloc[r]))
    ws.freeze_panes = "C2"
    ev = wb.create_sheet("evaluation")
    ev.column_dimensions["A"].width = 28
    top = 1
    for i, name in enumerate(preds):
        top = eval_block(ev, name, L(3 + i), n, top)

    sg = wb.create_sheet("subgroups")
    header(sg, 1, ["Group", "Rows", "No-show rate"] + [f"AUC {k}" for k in preds])
    groups = {f"Gender {g}": test["Gender"].to_numpy() == g for g in sorted(test["Gender"].unique())}
    groups.update({f"Age {a}": test["AgeBand"].astype(str).to_numpy() == a
                   for a in sorted(test["AgeBand"].astype(str).unique())})
    for r, (g, m) in enumerate(groups.items(), 2):
        sg.cell(r, 1, g); sg.cell(r, 2, int(m.sum())); sg.cell(r, 3, float(y[m].mean()))
        for c, p in enumerate(preds.values(), 4):
            sg.cell(r, c, float(roc_auc_score(y[m], p[m])))
    sg.column_dimensions["A"].width = 18

    # R4 inputs (rates learned on the training period)
    rs = wb.create_sheet("r4_rates")
    rs.cell(1, 1, "Global show rate (train)").font = BOLD; rs.cell(1, 2, float(r4info["s_bar"]))
    rs.cell(3, 1, "Group").font = BOLD; rs.cell(3, 2, "Smoothed show rate (k=5)").font = BOLD
    r = 4
    for k, v in {**{f"Gender {a}": b for a, b in r4info["gender"].items()},
                 **{f"Age {a}": b for a, b in r4info["age"].items()}}.items():
        rs.cell(r, 1, k); rs.cell(r, 2, float(v)); r += 1
    rs.column_dimensions["A"].width = 28
    wb.save(DATA / "auc_calibration.xlsx")


def main():
    df, train, test, cut = load()
    print(f"train {len(train)}, test {len(test)}, cut {cut:%Y-%m-%d}")
    p_lr = logistic_report(train, test, cut)
    p_rf = forest_report(train, test, cut)
    p_r4, info = r4_predict(train, test)
    preds = {"LogisticRegression": p_lr, "RandomForest": p_rf, "R4_rules": p_r4}
    y = test["NoShow"]
    for k, p in preds.items():
        print(k, {a: round(b, 4) for a, b in metrics(y, p).items()})
    comparison_report(train, test, cut, preds, info)
    # model risks and outcomes for the simulation (simulation/clinic_sim.py --risks ...)
    out = test[["AppointmentID", "ScheduledDay", "AppointmentDay", "Age", "LeadDays", "PriorApptCount"]].copy()
    out["p"], out["y"] = p_rf, y.to_numpy()
    out["p_lr"], out["p_r4"] = p_lr, p_r4  # other risk sources, for comparison (p = random forest)
    out.to_csv(DATA / "risks_random_forest.csv", index=False)


if __name__ == "__main__":
    main()
