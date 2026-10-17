"""Serves docs/web-designs and appends the demo bookings to a separate Excel file.

The static pages cannot write files, so they POST each booking here. The main data file
(ml/data/appointments_clean.xlsx) is never touched; rows go to output/web_randevulari.xlsx.

    python simulation/web_designs/serve.py        # then open http://127.0.0.1:8765/design1.html
"""

import json
import threading
from datetime import datetime
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
PAGES = ROOT / "docs" / "web-designs"
XLSX = ROOT / "output" / "web_randevulari.xlsx"
PORT = 8765
HEADERS = [
    "Kayıt zamanı",
    "Randevu no",
    "Ad soyad",
    "Bilinen hastalık",
    "Doğum tarihi",
    "Mahalle",
    "Randevu günü",
    "Saat",
    "Tür",
    "Risk (%)",
    "Durum",
]
LOCK = threading.Lock()


def _open():
    if XLSX.exists():
        wb = load_workbook(XLSX)
        return wb, wb.active
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Randevular"
    ws.append(HEADERS)
    for col, w in zip("ABCDEFGHIJK", (19, 12, 24, 24, 13, 18, 13, 7, 8, 9, 10), strict=True):
        ws.column_dimensions[col].width = w
    return wb, ws


def add_row(d: dict) -> None:
    with LOCK:
        wb, ws = _open()
        who = d.get("who") or {}
        ws.append(
            [
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                d["no"],
                who.get("name", ""),
                who.get("ill", "") or "Yok",
                who.get("birth", ""),
                who.get("hood", ""),
                d["iso"],
                d["time"],
                "Ek randevu" if d.get("kind") == "extra" else "Normal",
                d.get("risk", ""),
                "Aktif",
            ]
        )
        wb.save(XLSX)


def cancel_row(no: str) -> None:
    with LOCK:
        wb, ws = _open()
        for row in ws.iter_rows(min_row=2):
            if row[1].value == no:
                row[10].value = "İptal"
        wb.save(XLSX)


class Handler(SimpleHTTPRequestHandler):
    def _cors(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def do_OPTIONS(self):
        self.send_response(204)
        self._cors()
        self.end_headers()

    def do_POST(self):
        try:
            data = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))))
            {"/kayit": add_row, "/iptal": lambda d: cancel_row(d["no"])}[self.path](data)
            code = 200
        except PermissionError:  # Excel has the file open
            code = 503
        except Exception:
            code = 400
        self.send_response(code)
        self._cors()
        self.end_headers()
        self.wfile.write(b"{}")


if __name__ == "__main__":
    handler = partial(Handler, directory=str(PAGES))
    print(f"Sayfalar: http://127.0.0.1:{PORT}/design1.html  |  Excel: {XLSX}")
    ThreadingHTTPServer(("127.0.0.1", PORT), handler).serve_forever()
